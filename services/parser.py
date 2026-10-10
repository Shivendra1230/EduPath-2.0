"""Safe text extraction helpers for EduPath learning materials."""
import io
from pathlib import Path
import fitz
from docx import Document


def extract_document(uploaded_file) -> list[dict]:
    """Return source-aware sections with page/slide metadata."""
    name = uploaded_file.name
    suffix = Path(name).suffix.lower()
    data = uploaded_file.getvalue()
    sections = []
    if suffix == ".pdf":
        with fitz.open(stream=data, filetype="pdf") as doc:
            for i, page in enumerate(doc, 1):
                text = page.get_text("text").strip()
                if text:
                    sections.append({"text": text, "locator": f"page {i}", "source": name})
    elif suffix == ".docx":
        doc = Document(io.BytesIO(data))
        text = "\n".join(p.text for p in doc.paragraphs if p.text.strip())
        if text: sections.append({"text": text, "locator": "document", "source": name})
        for i, table in enumerate(doc.tables, 1):
            rows = [" | ".join(c.text.strip() for c in row.cells) for row in table.rows]
            if rows: sections.append({"text": "\n".join(rows), "locator": f"table {i}", "source": name})
    elif suffix == ".pptx":
        from pptx import Presentation
        prs = Presentation(io.BytesIO(data))
        for i, slide in enumerate(prs.slides, 1):
            bits=[]
            for shape in slide.shapes:
                if getattr(shape, "has_text_frame", False) and shape.text.strip(): bits.append(shape.text.strip())
                if getattr(shape, "has_table", False):
                    bits.extend(" | ".join(c.text.strip() for c in row.cells) for row in shape.table.rows)
        
            if bits: sections.append({"text": "\n".join(bits), "locator": f"slide {i}", "source": name})
    else:
        text = data.decode("utf-8", errors="ignore").strip()
        if text: sections.append({"text": text, "locator": "document", "source": name})
    return sections


def extract_text(uploaded_file) -> str:
    return "\n\n".join(s["text"] for s in extract_document(uploaded_file))


def clean_text(text: str) -> str:
    return "\n".join(line.strip() for line in text.splitlines() if line.strip())
